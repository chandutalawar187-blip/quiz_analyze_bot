"""Excel export tests."""

import unittest
from io import BytesIO

from openpyxl import load_workbook

from quizcord_bot.database import (
    QuizParticipant,
    QuizSession,
    QuizSessionSummary,
    StudentAnalytics,
    StudentQuizAnalytics,
    StudentQuizRecord,
)
from quizcord_bot.exports import (
    analytics_workbook,
    quiz_session_workbook,
    student_analytics_workbook,
)


class AnalyticsExportTests(unittest.TestCase):
    def test_exports_points_and_attendance_to_styled_workbook(self) -> None:
        student = StudentAnalytics(
            discord_id="12345",
            usn="=1+1",
            student_name="Student, One",
            quizzes_attended=2,
            quizzes_with_points=2,
            total_points=1800,
            average_points_per_scored_quiz=900,
        )

        workbook = load_workbook(BytesIO(analytics_workbook([student])))
        worksheet = workbook["Quizcord Analytics"]

        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            [
                "Discord ID",
                "Student Name",
                "USN",
                "Quizcord Quizzes Attended",
                "Quizzes With Points",
                "Total Quizcord Points",
                "Average Points Per Scored Quiz",
            ],
        )
        self.assertEqual(worksheet["B2"].value, "Student, One")
        self.assertEqual(worksheet["C2"].value, "'=1+1")
        self.assertEqual(
            [worksheet.cell(row=2, column=column).value for column in range(4, 8)],
            [2, 2, 1800, 900],
        )
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.auto_filter.ref, "A1:G2")
        self.assertEqual(worksheet["A1"].font.bold, True)
        self.assertEqual(worksheet["G2"].number_format, "0.00")

    def test_exports_complete_quiz_roster_including_unmapped_usns(self) -> None:
        session = QuizSession(
            session_id="session-1",
            guild_id=1,
            channel_id=2,
            status="completed",
            question_count=5,
            started_at="2026-10-17T15:00:00+00:00",
            completed_at="2026-10-17T15:30:00+00:00",
            quizcord_quiz_id="quiz-123",
            admin_user_id=3,
            admin_dm_channel_id=4,
            admin_dm_message_id=5,
            eligible_student_count=10,
        )
        summary = QuizSessionSummary(
            session=session,
            participants=(
                QuizParticipant("ABC123", 1, 900, "12345"),
                QuizParticipant("=UNMAPPED", 2, 700, None),
            ),
            total_points=1600,
            average_points=800.0,
            question_count=5,
        )

        workbook = load_workbook(BytesIO(quiz_session_workbook(summary)))

        self.assertEqual(workbook["Quiz Summary"]["B2"].value, "quiz-123")
        self.assertEqual(workbook["Quiz Summary"]["B5"].value, "2 / 10")
        self.assertEqual(workbook["Quiz Summary"]["B6"].value, 0.2)
        self.assertIn(
            "regardless of server permissions",
            workbook["Quiz Summary"]["B8"].value,
        )
        roster = workbook["Participants"]
        self.assertEqual(roster["B2"].value, "ABC123")
        self.assertEqual(roster["B3"].value, "'=UNMAPPED")
        self.assertEqual(roster["C3"].value, 700)

    def test_exports_student_summary_and_datewise_quiz_history(self) -> None:
        analytics = StudentQuizAnalytics(
            usn="ABC123",
            student_name="Student One",
            quizzes_attended=1,
            completed_quizzes=2,
            attendance_percentage=50.0,
            total_points=900,
            average_points=900.0,
            history=(
                StudentQuizRecord(
                    session_id="session-1",
                    channel_id=2,
                    quizcord_quiz_id="quiz-1",
                    status="completed",
                    started_at="2026-10-17T10:00:00+00:00",
                    completed_at="2026-10-17T10:10:00+00:00",
                    rank=1,
                    points=900,
                ),
                StudentQuizRecord(
                    session_id="import:quiz-2",
                    channel_id=None,
                    quizcord_quiz_id="=quiz-2",
                    status="historical",
                    started_at=None,
                    completed_at=None,
                    rank=0,
                    points=0,
                ),
            ),
        )

        workbook = load_workbook(BytesIO(student_analytics_workbook(analytics)))

        summary = workbook["Student Summary"]
        self.assertEqual(summary["B2"].value, "Student One")
        self.assertEqual(summary["B3"].value, "ABC123")
        self.assertEqual(summary["B6"].value, 0.5)
        self.assertEqual(summary["B7"].value, 900)
        history = workbook["Quiz History"]
        self.assertEqual(history["A2"].value, "2026-10-17T10:00:00+00:00")
        self.assertEqual(history["B2"].value, "quiz-1")
        self.assertEqual(history["C2"].value, "Completed")
        self.assertEqual(history["A3"].value, "Date unavailable")
        self.assertEqual(history["B3"].value, "'=quiz-2")
        self.assertEqual(history["D3"].value, "N/A")


if __name__ == "__main__":
    unittest.main()
