"""Excel export helpers for private student analytics reports."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from quizcord_bot.database import (
    QuizSessionSummary,
    StudentAnalytics,
    StudentQuizAnalytics,
)


def _safe_spreadsheet_text(value: str) -> str:
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


def analytics_workbook(report: list[StudentAnalytics]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Quizcord Analytics"
    headers = (
        "Discord ID",
        "Student Name",
        "USN",
        "Quizcord Quizzes Attended",
        "Quizzes With Points",
        "Total Quizcord Points",
        "Average Points Per Scored Quiz",
    )
    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24527A")
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for student in report:
        worksheet.append(
            [
                _safe_spreadsheet_text(student.discord_id),
                _safe_spreadsheet_text(student.student_name),
                _safe_spreadsheet_text(student.usn),
                student.quizzes_attended,
                student.quizzes_with_points,
                student.total_points,
                student.average_points_per_scored_quiz,
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32
    column_widths = (18, 28, 20, 28, 22, 24, 34)
    for index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows(min_row=2, min_col=7, max_col=7):
        row[0].number_format = "0.00"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def student_analytics_workbook(analytics: StudentQuizAnalytics) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Student Summary"
    overview.append(("Metric", "Value"))
    overview.append(("Student Name", _safe_spreadsheet_text(analytics.student_name)))
    overview.append(("USN", _safe_spreadsheet_text(analytics.usn)))
    overview.append(("Quizzes Attended", analytics.quizzes_attended))
    overview.append(("Completed Quizzes Tracked", analytics.completed_quizzes))
    overview.append(("Attendance Percentage", analytics.attendance_percentage / 100))
    overview.append(("Total Quizcord Points", analytics.total_points))
    overview.append(
        (
            "Average Points Per Attended Quiz",
            analytics.average_points if analytics.average_points is not None else "N/A",
        )
    )
    _style_header(overview)
    overview.freeze_panes = "A2"
    overview.column_dimensions["A"].width = 36
    overview.column_dimensions["B"].width = 32
    overview["B6"].number_format = "0.0%"

    history = workbook.create_sheet("Quiz History")
    history.append(("Date and Time", "Quiz ID", "Status", "Rank", "Points"))
    for record in analytics.history:
        history.append(
            [
                _safe_spreadsheet_text(record.started_at or "Date unavailable"),
                _safe_spreadsheet_text(record.quizcord_quiz_id or "Not provided"),
                _safe_spreadsheet_text(record.status.title()),
                record.rank if record.rank else "N/A",
                record.points,
            ]
        )
    _style_header(history)
    history.freeze_panes = "A2"
    history.auto_filter.ref = history.dimensions
    for column, width in (("A", 28), ("B", 24), ("C", 20), ("D", 12), ("E", 16)):
        history.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def quiz_session_workbook(summary: QuizSessionSummary) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Quiz Summary"
    overview.append(("Metric", "Value"))
    overview.append(("Quizcord Quiz ID", summary.session.quizcord_quiz_id or "Not provided"))
    overview.append(("Status", summary.session.status.title()))
    overview.append(("Leaderboard updates", summary.question_count))
    human_count = summary.session.eligible_student_count
    overview.append(
        (
            "Students attended"
            if summary.session.status == "completed"
            else "Students attending",
            f"{summary.attendance_count} / {human_count}"
            if human_count is not None
            else f"{summary.attendance_count} / unavailable",
        )
    )
    overview.append(
        (
            "Attendance percentage",
            summary.attendance_percentage / 100
            if summary.attendance_percentage is not None
            else "Unavailable",
        )
    )
    overview.append(
        (
            "Human server members",
            human_count if human_count is not None else "Unavailable",
        )
    )
    overview.append(
        (
            "Attendance counting",
            "Includes human members regardless of server permissions. Bot accounts are excluded.",
        )
    )
    _style_header(overview)
    overview.freeze_panes = "A2"
    overview.column_dimensions["A"].width = 30
    overview.column_dimensions["B"].width = 72
    overview["B6"].number_format = "0.0%"
    overview["B8"].alignment = Alignment(wrap_text=True, vertical="top")

    roster = workbook.create_sheet("Participants")
    roster.append(("Rank", "Student USN", "Final Quizcord Points"))
    for participant in summary.participants:
        roster.append(
            [
                participant.rank,
                _safe_spreadsheet_text(participant.usn),
                participant.points,
            ]
        )
    _style_header(roster)
    roster.freeze_panes = "A2"
    roster.auto_filter.ref = roster.dimensions
    roster.column_dimensions["A"].width = 12
    roster.column_dimensions["B"].width = 24
    roster.column_dimensions["C"].width = 28

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style_header(worksheet: Worksheet) -> None:
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24527A")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[1].height = 30
